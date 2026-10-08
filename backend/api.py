"""FastAPI and SSE interface for the research agent."""

from __future__ import annotations

import asyncio
import html
import json
import re
import unicodedata
from collections.abc import AsyncIterator
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from typing import Any
from uuid import uuid4

import markdown
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
from langgraph.types import Command
from pydantic import BaseModel, Field
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    HRFlowable,
    Paragraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sse_starlette.sse import EventSourceResponse
import reportlab

from backend.graph import agent_app


class ResearchRequest(BaseModel):
    """Request body for starting a research run."""

    query: str = Field(min_length=1, description="Question to research.")


class SourceReviewRequest(BaseModel):
    """Human-approved sources used to resume a paused research run."""

    sources: list[dict[str, str]]


class ExportRequest(BaseModel):
    """Markdown report and optional title to export as a brief."""

    report_markdown: str = Field(min_length=1)
    title: str = Field(default="Research Report", min_length=1, max_length=200)


app = FastAPI(title="Research Agent API", version="0.2.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

_pending_reviews: dict[str, asyncio.Future[list[dict[str, str]]]] = {}
_registered_pdf_fonts: tuple[str, str, str, str] | None = None


def _ensure_pdf_fonts() -> tuple[str, str, str, str]:
    """Register a Unicode TrueType font family, preferring Arial on Windows."""
    global _registered_pdf_fonts
    if _registered_pdf_fonts is not None:
        return _registered_pdf_fonts

    font_sets = [
        (
            Path("C:/Windows/Fonts/arial.ttf"),
            Path("C:/Windows/Fonts/arialbd.ttf"),
            Path("C:/Windows/Fonts/ariali.ttf"),
            Path("C:/Windows/Fonts/arialbi.ttf"),
        ),
        (
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-BoldOblique.ttf"),
        ),
    ]
    vera_dir = Path(reportlab.__file__).parent / "fonts"
    font_sets.append((vera_dir / "Vera.ttf", vera_dir / "VeraBd.ttf", vera_dir / "VeraIt.ttf", vera_dir / "VeraBI.ttf"))
    selected = next((font_set for font_set in font_sets if all(path.is_file() for path in font_set)), None)
    if selected is None:
        raise RuntimeError("No supported TrueType font family was found for PDF export.")

    names = (
        "ResearchSans",
        "ResearchSans-Bold",
        "ResearchSans-Italic",
        "ResearchSans-BoldItalic",
    )
    for name, path in zip(names, selected):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(path)))
    pdfmetrics.registerFontFamily(
        "ResearchSans",
        normal=names[0],
        bold=names[1],
        italic=names[2],
        boldItalic=names[3],
    )
    _registered_pdf_fonts = names
    return names


def clean_markdown_text(md_text: str) -> str:
    """Normalize Markdown whitespace, punctuation, and control characters."""
    cleaned = md_text.replace("\xa0", " ").replace("\u202f", " ")
    cleaned = cleaned.translate(str.maketrans({
        "‐": "-", "‑": "-", "‒": "-", "–": "-", "—": "-", "―": "-", "−": "-",
        "‘": "'", "’": "'", "‚": "'", "‛": "'", "“": '"', "”": '"', "„": '"', "‟": '"',
        "…": "...",
    }))
    cleaned = unicodedata.normalize("NFKC", cleaned)
    return "".join(
        character for character in cleaned
        if character in "\n\r\t" or unicodedata.category(character) not in {"Cc", "Cf", "Cs"}
    )


def _clean_pdf_text(value: str, font_name: str) -> str:
    """Normalize whitespace and punctuation and replace glyphs absent in the PDF font."""
    value = clean_markdown_text(value)
    font_face = pdfmetrics.getFont(font_name).face
    char_to_glyph = getattr(font_face, "charToGlyph", {})
    cleaned: list[str] = []
    for character in value:
        if character in "\n\r\t" or not char_to_glyph or ord(character) in char_to_glyph:
            cleaned.append(character)
            continue
        ascii_fallback = "".join(
            part for part in unicodedata.normalize("NFKD", character)
            if part.isascii() and part.isprintable() and not unicodedata.combining(part)
        )
        cleaned.append(ascii_fallback or "?")
    return "".join(cleaned)


@app.get("/health")
async def health() -> dict[str, str]:
    """Return service health status."""
    return {"status": "ok"}


@app.post("/api/export/pdf")
def export_pdf(request: ExportRequest) -> Response:
    """Render structured Markdown content as a styled PDF attachment."""
    font_names = _ensure_pdf_fonts()
    cleaned_markdown = _clean_pdf_text(clean_markdown_text(request.report_markdown), font_names[0])
    rendered_html = markdown.markdown(
        cleaned_markdown,
        extensions=["tables", "fenced_code", "nl2br"],
        output_format="html5",
    )

    sample_styles = getSampleStyleSheet()
    pdf_title = _clean_pdf_text(request.title, font_names[0])
    title_style = ParagraphStyle(
        "ResearchTitle", parent=sample_styles["Title"], fontName=font_names[1],
        fontSize=23, leading=28, textColor=colors.HexColor("#0f172a"),
        alignment=TA_LEFT, spaceAfter=14,
    )
    heading_styles = {
        1: ParagraphStyle("ResearchH1", parent=sample_styles["Heading1"], fontName=font_names[1], fontSize=18, leading=22, textColor=colors.HexColor("#0f172a"), spaceBefore=15, spaceAfter=7, keepWithNext=True),
        2: ParagraphStyle("ResearchH2", parent=sample_styles["Heading2"], fontName=font_names[1], fontSize=14, leading=18, textColor=colors.HexColor("#1e293b"), spaceBefore=12, spaceAfter=6, keepWithNext=True),
        3: ParagraphStyle("ResearchH3", parent=sample_styles["Heading3"], fontName=font_names[1], fontSize=11, leading=14, textColor=colors.HexColor("#334155"), spaceBefore=9, spaceAfter=4, keepWithNext=True),
    }
    body_style = ParagraphStyle(
        "ResearchBody", parent=sample_styles["BodyText"], fontName=font_names[0],
        fontSize=9.5, leading=14, textColor=colors.HexColor("#334155"), spaceAfter=7,
    )
    quote_style = ParagraphStyle(
        "ResearchQuote", parent=body_style, leftIndent=14,
        borderColor=colors.HexColor("#94a3b8"), borderWidth=2, borderPadding=6,
        textColor=colors.HexColor("#475569"),
    )
    code_style = ParagraphStyle(
        "ResearchCode", parent=body_style, fontName=font_names[0], fontSize=8,
        leading=10, backColor=colors.HexColor("#f1f5f9"), borderPadding=7,
    )
    story: list[Any] = [
        Paragraph(html.escape(pdf_title), title_style),
        HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cbd5e1")),
        Spacer(1, 12),
    ]

    paragraph_styles = {"p": body_style, "blockquote": quote_style, "li": body_style}

    class ReportHtmlParser(HTMLParser):
        """Translate Markdown HTML into ReportLab paragraphs and wrapped tables."""

        def __init__(self) -> None:
            super().__init__(convert_charrefs=True)
            self.parts: list[str] = []
            self.block_tag: str | None = None
            self.in_cell = False
            self.cell_is_header = False
            self.cell_parts: list[str] = []
            self.rows: list[list[Paragraph]] = []
            self.row: list[Paragraph] = []
            self.list_stack: list[tuple[str, int]] = []
            self.open_links = 0

        def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
            attributes = dict(attrs)
            if tag == "table":
                self.rows = []
            elif tag == "tr":
                self.row = []
            elif tag in {"th", "td"}:
                self.in_cell = True
                self.cell_is_header = tag == "th"
                self.cell_parts = ["<b>"] if self.cell_is_header else []
            elif tag in {"ul", "ol"}:
                self.list_stack.append((tag, 0))
            elif tag in {"p", "blockquote", "li", "h1", "h2", "h3", "h4", "h5", "h6", "pre"} and not self.in_cell:
                self.block_tag = tag
                self.parts = []
                if tag == "li" and self.list_stack:
                    list_kind, count = self.list_stack[-1]
                    count += 1
                    self.list_stack[-1] = (list_kind, count)
                    marker = "&#8226; " if list_kind == "ul" else f"{count}. "
                    self.parts.append(marker)
            elif tag in {"strong", "b"}:
                self._append("<b>")
            elif tag in {"em", "i"}:
                self._append("<i>")
            elif tag == "code":
                self._append(f'<font name="{font_names[0]}">')
            elif tag == "a":
                href = attributes.get("href") or ""
                if href.startswith(("http://", "https://", "mailto:")):
                    self._append(f'<link href="{html.escape(href, quote=True)}" color="#2563eb">')
                    self.open_links += 1
            elif tag == "br":
                self._append("<br/>")
            elif tag == "hr":
                story.append(HRFlowable(width="100%", thickness=0.6, color=colors.HexColor("#cbd5e1"), spaceBefore=5, spaceAfter=10))

        def handle_endtag(self, tag: str) -> None:
            if tag in {"th", "td"} and self.in_cell:
                if self.cell_is_header:
                    self.cell_parts.append("</b>")
                content = "".join(self.cell_parts).strip() or " "
                style = table_header_style if self.cell_is_header else table_cell_style
                self.row.append(Paragraph(content, style))
                self.in_cell = False
            elif tag == "tr":
                if self.row:
                    self.rows.append(self.row)
                    self.row = []
            elif tag == "table":
                self._flush_block()
                if self.rows:
                    column_count = max(map(len, self.rows))
                    width = (A4[0] - 40 * mm) / column_count
                    padded_rows = [row + [Paragraph(" ", table_cell_style)] * (column_count - len(row)) for row in self.rows]
                    table = Table(padded_rows, colWidths=[width] * column_count, repeatRows=1, splitByRow=1, hAlign="LEFT")
                    table.setStyle(TableStyle([
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 6),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                        ("TOPPADDING", (0, 0), (-1, -1), 5),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ]))
                    story.extend([table, Spacer(1, 9)])
                self.rows = []
            elif tag in {"ul", "ol"}:
                if self.list_stack:
                    self.list_stack.pop()
            elif tag in {"strong", "b"}:
                self._append("</b>")
            elif tag in {"em", "i"}:
                self._append("</i>")
            elif tag == "code":
                self._append("</font>")
            elif tag == "a":
                if self.open_links:
                    self._append("</link>")
                    self.open_links -= 1
            elif tag in {"p", "blockquote", "li", "h1", "h2", "h3", "h4", "h5", "h6", "pre"} and not self.in_cell:
                self._flush_block()

        def handle_data(self, data: str) -> None:
            escaped = html.escape(data, quote=False)
            self._append(escaped)

        def _append(self, value: str) -> None:
            if self.in_cell:
                self.cell_parts.append(value)
            elif self.block_tag:
                self.parts.append(value)

        def _flush_block(self) -> None:
            if self.block_tag is None:
                return
            content = "".join(self.parts).strip()
            tag = self.block_tag
            self.block_tag = None
            self.parts = []
            if not content:
                return
            if tag.startswith("h") and tag[1:].isdigit():
                level = min(max(int(tag[1:]), 1), 3)
                story.append(Paragraph(content, heading_styles[level]))
            elif tag == "pre":
                story.append(Preformatted(html.unescape(content), code_style, maxLineLength=95))
            else:
                story.append(Paragraph(content, paragraph_styles.get(tag, body_style)))

    table_cell_style = ParagraphStyle("ResearchTableCell", parent=body_style, fontSize=8, leading=11, spaceAfter=0, wordWrap="CJK")
    table_header_style = ParagraphStyle("ResearchTableHeader", parent=table_cell_style, fontName=font_names[1], textColor=colors.HexColor("#0f172a"))
    parser = ReportHtmlParser()
    parser.feed(rendered_html)
    parser.close()

    def draw_page_number(canvas: Any, document: Any) -> None:
        canvas.saveState()
        canvas.setFont(font_names[0], 8)
        canvas.setFillColor(colors.HexColor("#64748b"))
        canvas.drawRightString(A4[0] - 20 * mm, 12 * mm, str(document.page))
        canvas.restoreState()

    try:
        output = BytesIO()
        pdf = SimpleDocTemplate(
            output, pagesize=A4, rightMargin=20 * mm, leftMargin=20 * mm,
            topMargin=20 * mm, bottomMargin=20 * mm, title=pdf_title,
        )
        pdf.build(story, onFirstPage=draw_page_number, onLaterPages=draw_page_number)
        pdf_bytes = output.getvalue()
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Could not generate PDF: {exc}") from exc
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="research_report.pdf"'},
    )


@app.post("/api/export/docx")
def export_docx(request: ExportRequest) -> Response:
    """Convert parsed Markdown headings, paragraphs, lists, and tables to DOCX."""
    cleaned_markdown = clean_markdown_text(request.report_markdown)
    rendered_html = markdown.markdown(
        cleaned_markdown,
        extensions=["tables", "fenced_code", "nl2br"],
        output_format="html5",
    )
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(51, 65, 85)
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.color.rgb = RGBColor(15, 23, 42)

    first_content_line = next((line.strip() for line in cleaned_markdown.splitlines() if line.strip()), "")
    markdown_title = re.sub(r"^#\s+", "", first_content_line)
    if not first_content_line.startswith("# ") or markdown_title.casefold() != request.title.casefold():
        title_paragraph = document.add_heading(request.title, level=0)
        title_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    class DocxHtmlParser(HTMLParser):
        """Build native Word paragraphs, formatted runs, and grid tables."""

        def __init__(self) -> None:
            super().__init__(convert_charrefs=True)
            self.paragraph = None
            self.bold = 0
            self.italic = 0
            self.code = 0
            self.lists: list[str] = []
            self.rows: list[list[tuple[list[tuple[str, bool, bool, bool]], bool]]] = []
            self.row: list[tuple[list[tuple[str, bool, bool, bool]], bool]] = []
            self.cell_runs: list[tuple[str, bool, bool, bool]] | None = None
            self.cell_header = False
            self.in_pre = False

        def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
            if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
                self.paragraph = document.add_heading("", level=int(tag[1]))
            elif tag == "p" and self.paragraph is None and self.cell_runs is None:
                self.paragraph = document.add_paragraph()
            elif tag in {"ul", "ol"}:
                self.lists.append(tag)
            elif tag == "li" and self.cell_runs is None:
                style = "List Number" if self.lists and self.lists[-1] == "ol" else "List Bullet"
                self.paragraph = document.add_paragraph(style=style)
            elif tag == "table":
                self.rows = []
            elif tag == "tr":
                self.row = []
            elif tag in {"td", "th"}:
                self.cell_runs = []
                self.cell_header = tag == "th"
            elif tag in {"strong", "b"}:
                self.bold += 1
            elif tag in {"em", "i"}:
                self.italic += 1
            elif tag == "code":
                self.code += 1
            elif tag == "pre":
                self.in_pre = True
                self.paragraph = document.add_paragraph()
            elif tag == "br":
                if self.cell_runs is not None:
                    self._append_cell_text("\n")
                elif self.paragraph is not None:
                    self.paragraph.add_run().add_break()

        def handle_endtag(self, tag: str) -> None:
            if tag in {"td", "th"} and self.cell_runs is not None:
                self.row.append((self.cell_runs, self.cell_header))
                self.cell_runs = None
            elif tag == "tr":
                if self.row:
                    self.rows.append(self.row)
                    self.row = []
            elif tag == "table":
                self._write_table()
            elif tag in {"ul", "ol"}:
                if self.lists:
                    self.lists.pop()
            elif tag in {"strong", "b"}:
                self.bold = max(0, self.bold - 1)
            elif tag in {"em", "i"}:
                self.italic = max(0, self.italic - 1)
            elif tag == "code":
                self.code = max(0, self.code - 1)
            elif tag == "pre":
                self.in_pre = False
                self.paragraph = None
            elif tag in {"p", "li", "h1", "h2", "h3", "h4", "h5", "h6"}:
                self.paragraph = None

        def handle_data(self, data: str) -> None:
            if self.cell_runs is not None:
                self._append_cell_text(data)
                return
            if not data.strip() and self.paragraph is None:
                return
            if self.paragraph is None:
                self.paragraph = document.add_paragraph()
            run = self.paragraph.add_run(data)
            run.bold = self.bold > 0
            run.italic = self.italic > 0
            if self.code or self.in_pre:
                run.font.name = "Consolas"

        def _append_cell_text(self, text: str) -> None:
            assert self.cell_runs is not None
            self.cell_runs.append((text, self.bold > 0 or self.cell_header, self.italic > 0, self.code > 0))

        def _write_table(self) -> None:
            if not self.rows:
                return
            column_count = max(len(row) for row in self.rows)
            table = document.add_table(rows=len(self.rows), cols=column_count)
            table.style = "Table Grid"
            table.autofit = True
            for row_index, row_data in enumerate(self.rows):
                for column_index, (runs, is_header) in enumerate(row_data):
                    cell = table.cell(row_index, column_index)
                    cell.text = ""
                    cell_properties = cell._tc.get_or_add_tcPr()
                    margins = OxmlElement("w:tcMar")
                    for side in ("top", "start", "bottom", "end"):
                        margin = OxmlElement(f"w:{side}")
                        margin.set(qn("w:w"), "90")
                        margin.set(qn("w:type"), "dxa")
                        margins.append(margin)
                    cell_properties.append(margins)
                    if is_header:
                        shading = OxmlElement("w:shd")
                        shading.set(qn("w:fill"), "F1F5F9")
                        cell_properties.append(shading)
                    target = cell.paragraphs[0]
                    for text, bold, italic, code in runs:
                        run = target.add_run(text)
                        run.bold = bold or is_header
                        run.italic = italic
                        if code:
                            run.font.name = "Consolas"
                    if not runs:
                        target.add_run(" ")
            self.rows = []

    docx_parser = DocxHtmlParser()
    docx_parser.feed(rendered_html)
    docx_parser.close()

    buffer = BytesIO()
    document.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": 'attachment; filename="research_report.docx"'},
    )


@app.post("/api/research/{run_id}/sources")
async def approve_sources(run_id: str, request: SourceReviewRequest) -> dict[str, str]:
    """Submit source approval for an active SSE research run."""
    future = _pending_reviews.get(run_id)
    if future is None or future.done():
        raise HTTPException(status_code=404, detail="No active source review for this run.")
    future.set_result(request.sources)
    return {"status": "accepted", "run_id": run_id}


@app.post("/api/research/stream")
async def stream_research(request: ResearchRequest) -> EventSourceResponse:
    """Stream graph progress and pause for human source approval."""
    run_id = str(uuid4())

    async def event_generator() -> AsyncIterator[dict[str, str]]:
        config = {"configurable": {"thread_id": run_id}}
        review_future: asyncio.Future[list[dict[str, str]]] | None = None
        try:
            async for update in agent_app.astream(
                {"query": request.query}, config=config, stream_mode="updates"
            ):
                interrupt_events = update.get("__interrupt__", ())
                if interrupt_events:
                    interrupt_value: dict[str, Any] = interrupt_events[0].value
                    review_future = asyncio.get_running_loop().create_future()
                    _pending_reviews[run_id] = review_future
                    yield {
                        "event": "source_review",
                        "data": json.dumps(
                            {
                                "run_id": run_id,
                                "query": interrupt_value.get("query", request.query),
                                "search_results": interrupt_value.get("search_results", []),
                            }
                        ),
                    }
                    selected_sources = await review_future
                    _pending_reviews.pop(run_id, None)

                    async for resumed in agent_app.astream(
                        Command(resume=selected_sources),
                        config=config,
                        stream_mode="updates",
                    ):
                        for node_name, node_state in resumed.items():
                            if node_name != "__interrupt__":
                                yield {
                                    "event": "node_complete",
                                    "data": json.dumps(
                                        {"node": node_name, "state": node_state},
                                        default=str,
                                    ),
                                }
                    break

                for node_name, node_state in update.items():
                    yield {
                        "event": "node_complete",
                        "data": json.dumps(
                            {"node": node_name, "state": node_state}, default=str
                        ),
                    }
            yield {"event": "done", "data": json.dumps({"run_id": run_id})}
        except asyncio.CancelledError:
            raise
        except Exception as exc:
            yield {
                "event": "error",
                "data": json.dumps({"message": str(exc), "run_id": run_id}),
            }
        finally:
            _pending_reviews.pop(run_id, None)
            if review_future is not None and not review_future.done():
                review_future.cancel()

    return EventSourceResponse(event_generator())
